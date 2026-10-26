"""Quarterly capital account statements, whatever shape a fund manager sends them in.

The parser reads labels, not layouts. It turns any statement (PDF, Excel, CSV, text) into a sequence of cells,
finds the cells whose text starts with a known label (config/vocabulary.yaml, English and German), and reads
the value from the rest of that cell or the cells after it. Numbers are read in the document's own convention
(1,234.56 or 1.234,56), amounts are scaled when the document says they are in thousands, and dates are read in
whichever of the usual formats they come in.

A statement that cannot be read, or reads but breaks the contract (config/contracts/gp_statement.yaml), never
reaches the warehouse: it goes to quarantine with the rule it broke.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import re
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from ..config import target_funds, vocabulary

AMOUNT_FIELDS = ("commitment", "paid_in", "distributed", "nav", "unfunded")
DATE_FIELDS = ("period_end", "statement_date")
MONTHS = {
    m: i + 1
    for i, names in enumerate(
        [
            ("january", "jan", "januar", "jänner"),
            ("february", "feb", "februar"),
            ("march", "mar", "märz", "maerz"),
            ("april", "apr"),
            ("may", "mai"),
            ("june", "jun", "juni"),
            ("july", "jul", "juli"),
            ("august", "aug"),
            ("september", "sep", "sept"),
            ("october", "oct", "oktober", "okt"),
            ("november", "nov"),
            ("december", "dec", "dezember", "dez"),
        ]
    )
    for m in names
}
GERMAN_HINTS = (
    "stichtag",
    "zeichnungsbetrag",
    "eingezahlt",
    "ausschüttung",
    "währung",
    "anleger",
    "kapitalkonto",
    "erstellt am",
)


class ParseError(Exception):
    def __init__(self, rule: str, detail: str):
        super().__init__(f"{rule}: {detail}")
        self.rule = rule
        self.detail = detail


@dataclass
class Statement:
    fund: str | None = None
    period_end: dt.date | None = None
    statement_date: dt.date | None = None
    currency: str | None = None
    commitment: Decimal | None = None
    paid_in: Decimal | None = None
    distributed: Decimal | None = None
    nav: Decimal | None = None
    unfunded: Decimal | None = None
    unit_scale: int = 1
    parser: str = ""
    notes: list[str] = field(default_factory=list)

    def as_dict(self) -> dict[str, Any]:
        return {
            k: getattr(self, k)
            for k in ("fund", "period_end", "statement_date", "currency", *AMOUNT_FIELDS, "unit_scale")
        }


# --------------------------------------------------------------------------- cells from documents


def cells_from_file(path: Path) -> tuple[list[str], str]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _pdf_cells(path), "pdf"
    if suffix in (".xlsx", ".xlsm"):
        return _xlsx_cells(path), "xlsx"
    if suffix == ".csv":
        return _csv_cells(path.read_text(encoding="utf-8-sig")), "csv"
    if suffix in (".txt", ".md"):
        return _text_cells(path.read_text(encoding="utf-8")), "text"
    raise ParseError("unsupported_format", f"{path.suffix} is not a statement format the platform reads")


def _pdf_cells(path: Path) -> list[str]:
    from pypdf import PdfReader

    try:
        reader = PdfReader(str(path))
        text = "\n".join((p.extract_text() or "") for p in reader.pages)
    except Exception as e:  # a damaged PDF is a quarantine reason, not a crash
        raise ParseError("unreadable_pdf", str(e)[:200]) from e
    if not text.strip():
        raise ParseError("no_text_layer", "the PDF has no text layer (a scan); ask the manager for the original")
    return _text_cells(text)


def _text_cells(text: str) -> list[str]:
    out: list[str] = []
    for line in text.splitlines():
        line = line.strip().strip("|").strip()
        if not line or set(line) <= set("-=|: "):
            continue
        # table rows in text or Markdown: "| Label | Value |"
        parts = [p.strip() for p in re.split(r"\s*\|\s*|\t|\s{3,}", line) if p.strip()]
        out.extend(parts or [line])
    return out


def _xlsx_cells(path: Path) -> list[str]:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=True)
    out: list[str] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for v in row:
                if v is None or (isinstance(v, str) and not v.strip()):
                    continue
                if isinstance(v, dt.datetime):
                    out.append(v.date().isoformat())
                elif isinstance(v, dt.date):
                    out.append(v.isoformat())
                elif isinstance(v, float | int) and not isinstance(v, bool):
                    # six decimals, so a German reading never mistakes 123.456 for a thousands separator
                    out.append(f"{v:.6f}")
                else:
                    out.append(str(v).strip())
    return out


def _csv_cells(text: str) -> list[str]:
    sample = text[:2048]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delim = dialect.delimiter
    except csv.Error:
        delim = ";" if sample.count(";") > sample.count(",") else ","
    rows = [r for r in csv.reader(io.StringIO(text), delimiter=delim) if any(c.strip() for c in r)]
    if len(rows) < 2:
        raise ParseError("empty_csv", "a header and at least one row are needed")
    header, values = rows[0], rows[1]
    if len(rows) > 2:
        raise ParseError("several_statements", f"{len(rows) - 1} rows: one statement per file is expected")
    out: list[str] = []
    for h, v in zip(header, values, strict=False):
        out.extend([h.strip(), v.strip()])
    return out


# --------------------------------------------------------------------------- reading values


def _norm(s: str) -> str:
    s = s.lower().replace("_", " ").replace("’", "'")
    return re.sub(r"\s+", " ", s).strip(" :.\t")


def _labels() -> list[tuple[str, str]]:
    pairs = [(field_name, _norm(syn)) for field_name, syns in vocabulary()["fields"].items() for syn in syns]
    return sorted(pairs, key=lambda p: -len(p[1]))


def match_label(cell: str) -> tuple[str, str] | None:
    """(field, rest of the cell) when the cell starts with a known label."""
    n = _norm(cell)
    for field_name, syn in _labels():
        if n == syn or n.startswith(syn + " ") or n.startswith(syn + ":"):
            rest = cell.strip()[len(syn) :] if cell.strip().lower().startswith(syn) else _after_label(cell, syn)
            return field_name, rest.strip(" :\t")
    return None


def _after_label(cell: str, syn: str) -> str:
    # the label may differ in case, underscores or spacing: cut by the number of words
    words = cell.replace("_", " ").split()
    return " ".join(words[len(syn.split()) :])


def parse_number(raw: str, german: bool) -> Decimal | None:
    """A number written either way: 1,234,567.89 or 1.234.567,89, with currency signs, spaces or parentheses."""
    s = raw.strip()
    neg = (s.startswith("(") and s.endswith(")")) or s.startswith("-") or s.endswith("-")
    s = re.sub(r"[^\d.,]", "", s.replace("'", ""))
    if not re.search(r"\d", s):
        return None
    if "," in s and "." in s:
        dec = "," if s.rfind(",") > s.rfind(".") else "."
    elif "," in s:
        dec = "" if not german and re.fullmatch(r"\d{1,3}(,\d{3})+", s) else ","
    elif "." in s:
        dec = "" if (german and re.fullmatch(r"\d{1,3}(\.\d{3})+", s)) or s.count(".") > 1 else "."
    else:
        dec = ""
    thousands = {",": ".", ".": ",", "": ",."}[dec]
    for ch in thousands:
        s = s.replace(ch, "")
    s = s.replace(",", ".")
    try:
        v = Decimal(s)
    except InvalidOperation:
        return None
    return -v if neg else v


DATE_PATTERNS = [
    (re.compile(r"(\d{4})-(\d{2})-(\d{2})"), "iso"),
    (re.compile(r"(\d{1,2})\.(\d{1,2})\.(\d{4})"), "dmy_dot"),
    (re.compile(r"(\d{1,2})/(\d{1,2})/(\d{4})"), "slash"),
    (re.compile(r"([A-Za-zÄÖÜäöü]+)\.? (\d{1,2}),? (\d{4})"), "mdy_words"),
    (re.compile(r"(\d{1,2})\.? ([A-Za-zÄÖÜäöü]+)\.? (\d{4})"), "dmy_words"),
]


def parse_date(raw: str, prefer_us: bool | None = None) -> dt.date | None:
    for rx, kind in DATE_PATTERNS:
        m = rx.search(raw)
        if not m:
            continue
        try:
            if kind == "iso":
                return dt.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
            if kind == "dmy_dot":
                return dt.date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
            if kind == "slash":
                a, b, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
                if a > 12:
                    return dt.date(y, b, a)
                if b > 12:
                    return dt.date(y, a, b)
                return dt.date(y, a, b) if prefer_us else dt.date(y, b, a)
            if kind == "mdy_words":
                mon = MONTHS.get(m.group(1).lower())
                if mon:
                    return dt.date(int(m.group(3)), mon, int(m.group(2)))
            if kind == "dmy_words":
                mon = MONTHS.get(m.group(2).lower())
                if mon:
                    return dt.date(int(m.group(3)), mon, int(m.group(1)))
        except ValueError:
            continue
    return None


def _unit_scale(text: str) -> int:
    t = " " + _norm(text) + " "
    units = vocabulary()["units"]
    if any(re.search(r"(?<![a-z])" + re.escape(_norm(u)) + r"(?![a-z])", t) for u in units["millions"]):
        return 1_000_000
    if any(re.search(r"(?<![a-z])" + re.escape(_norm(u)) + r"(?![a-z])", t) for u in units["thousands"]):
        return 1000
    return 1


def _currency(text: str) -> str | None:
    up = text.upper()
    for cur in vocabulary()["currencies"]:
        if re.search(rf"(?<![A-Z]){cur}(?![A-Z])", up) or re.search(rf"(?<![A-Z])T{cur}(?![A-Z])", up):
            return cur
    if "€" in text:
        return "EUR"
    if "$" in text:
        return "USD"
    return None


def identify_fund(cells: list[str], stated: str | None) -> str | None:
    """The target fund a statement is about: by its name (without legal suffixes) anywhere in the document."""
    text = _norm(" ".join(cells))
    best = None
    for fid, f in target_funds().items():
        name = _norm(f["name"])
        if name in text or (stated and _norm(stated).startswith(name)):
            if best is None or len(name) > len(_norm(target_funds()[best]["name"])):
                best = fid
        elif re.search(rf"(?<![a-z0-9]){fid.lower()}(?![a-z0-9])", text):
            best = best or fid
    return best


def parse_cells(cells: list[str], parser: str) -> Statement:
    st = Statement(parser=parser)
    doc = " ".join(cells)
    german = sum(h in doc.lower() for h in GERMAN_HINTS) >= 2
    us_dates = bool(re.search(r"\b(0?[1-9]|1[0-2])/(29|30|31)/\d{4}\b", doc))
    found: dict[str, str] = {}
    for i, cell in enumerate(cells):
        m = match_label(cell)
        if not m:
            continue
        name, rest = m
        if name in found:
            continue
        candidates = [rest] if rest else []
        candidates += cells[i + 1 : i + 3]
        for cand in candidates:
            if name in AMOUNT_FIELDS:
                if parse_number(cand, german) is not None and not match_label(cand):
                    found[name] = cand
                    break
            elif name in DATE_FIELDS:
                if parse_date(cand, us_dates) is not None:
                    found[name] = cand
                    break
            elif name == "currency":
                if _currency(cand):
                    found[name] = cand
                    break
            elif name == "fund" and cand and not match_label(cand):
                found[name] = cand
                break
    st.unit_scale = _unit_scale(doc)
    for name in AMOUNT_FIELDS:
        if name in found:
            v = parse_number(found[name], german)
            if v is not None:
                setattr(st, name, (v * st.unit_scale).quantize(Decimal("0.01")))
    for name in DATE_FIELDS:
        if name in found:
            setattr(st, name, parse_date(found[name], us_dates))
    st.currency = _currency(found.get("currency", "")) or _currency(" ".join(found.get(f, "") for f in AMOUNT_FIELDS))
    st.fund = identify_fund(cells, found.get("fund"))
    if st.currency is None and st.fund:
        st.currency = target_funds()[st.fund]["currency"]
        st.notes.append("currency taken from the fund's master data: the statement does not state it")
    if st.unit_scale != 1:
        st.notes.append(f"amounts scaled by {st.unit_scale}: the statement says so")
    return st


def parse_statement(path: Path) -> Statement:
    cells, kind = cells_from_file(path)
    return parse_cells(cells, kind)


# --------------------------------------------------------------------------- contract


def _is_quarter_end(d: dt.date) -> bool:
    return d.month in (3, 6, 9, 12) and (d + dt.timedelta(days=1)).day == 1


def violations(st: Statement) -> list[tuple[str, str]]:
    """Rules of config/contracts/gp_statement.yaml the statement breaks, as (rule, detail)."""
    out: list[tuple[str, str]] = []
    missing = [
        f
        for f in ("fund", "period_end", "currency", "commitment", "paid_in", "distributed", "nav")
        if getattr(st, f) is None
    ]
    if missing:
        out.append(("required_fields", "missing: " + ", ".join(missing)))
        return out
    assert st.period_end and st.currency and st.commitment is not None and st.paid_in is not None
    assert st.distributed is not None and st.nav is not None
    if st.currency not in vocabulary()["currencies"]:
        out.append(("currency_known", st.currency))
    if not _is_quarter_end(st.period_end):
        out.append(("quarter_end", f"{st.period_end} is not the last day of a quarter"))
    if min(st.commitment, st.paid_in, st.distributed, st.nav) < 0:
        out.append(("amounts_not_negative", "an amount is negative"))
    if st.paid_in > st.commitment * Decimal("1.02"):
        out.append(("paid_in_within_commitment", f"paid in {st.paid_in} against a commitment of {st.commitment}"))
    if st.unfunded is not None and abs(st.commitment - st.paid_in - st.unfunded) > st.commitment * Decimal("0.005"):
        out.append(("unfunded_adds_up", f"{st.commitment} - {st.paid_in} != {st.unfunded}"))
    if st.statement_date is not None and st.statement_date < st.period_end:
        out.append(("statement_after_period", f"issued {st.statement_date} before the period end {st.period_end}"))
    return out
